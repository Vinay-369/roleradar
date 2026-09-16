import sys
import os
import time
import json
import fitz
from docx import Document
import io

sys.path.insert(0, os.path.abspath("backend"))

import copy
from app.modules.resume.models import (
    CandidateProfile,
    EvidenceUnit,
    WorkExperienceEntity,
    ProjectEntity,
    EducationEntity,
    ClaimType,
)
from app.modules.tailoring.strategy import (
    build_resume_strategy,
    compute_skill_priorities,
    compute_project_priorities,
    CareerStage,
    TemplateFamily,
)
from app.modules.tailoring.validation import (
    detect_fabricated_claims,
    detect_unsupported_metrics,
    detect_unsupported_action_verbs_and_scope,
    detect_sentence_fragments_and_truncation,
    detect_entity_boundary_violations,
    validate_protected_sections,
    validate_final_tailored_resume,
    validate_tailored_profile_truth_guard,
    compute_deterministic_skill_reorder,
)
from app.modules.tailoring.services import _audit_user_edits, _truth_guard_warning
from app.modules.intelligence.ats_score import compute_ats_score
from app.modules.intelligence.ats_readability_validator import evaluate_ats_and_readability
from app.modules.tailoring.export import (
    generate_pdf,
    generate_docx,
    verify_export_against_structured_resume,
    verify_ats_pdf_parseability,
)

print("="*70)
print("STARTING PHASE 6 COMPREHENSIVE AUDIT EXECUTION")
print("="*70)

# ============================================================
# 1. TRUTH GUARD ADVERSARIAL TEST SUITE (TG-01 through TG-08 + Semantic)
# ============================================================
print("\n--- 1. TRUTH GUARD ADVERSARIAL AUDIT ---")
candidate_skills = ["Python", "React", "MongoDB"]
orig_bullet_1 = "Built a web application using React."
orig_bullet_proj = "Built an e-commerce website."
orig_exp_bullet = "Assisted senior developers with frontend bug fixes."

# TG-01: Invented skill (AWS required in JD, candidate has Python, React, MongoDB)
tg01_proposed = "Experienced AWS developer building scalable cloud native microservices."
tg01_result = detect_fabricated_claims(orig_bullet_1, tg01_proposed, "Requires AWS, Docker, Kubernetes", candidate_skills)
print(f"TG-01 (Invented Skill 'AWS'): Blocked={bool('aws' in [x.lower() for x in tg01_result])}, Flagged={tg01_result}")

# TG-02: Invented experience (Fresher -> 3 years professional experience)
tg02_proposed = "Over 3 years of professional experience leading enterprise systems."
tg02_scope = detect_unsupported_action_verbs_and_scope("Student project intern.", tg02_proposed)
tg02_metrics = detect_unsupported_metrics("Student project intern.", tg02_proposed)
print(f"TG-02 (Invented Experience / Seniority): Scope Flagged={tg02_scope}, Metrics Flagged={tg02_metrics}")

# TG-03: Invented metric ("Improved performance by 45%" from unquantified project)
tg03_proposed = "Built an e-commerce website, improving performance by 45%."
tg03_metrics = detect_unsupported_metrics(orig_bullet_proj, tg03_proposed)
print(f"TG-03 (Invented Metric '45%'): Blocked={bool(tg03_metrics)}, Flagged={tg03_metrics}")

# TG-04: Invented certification (AWS Certified Developer added without source)
master_parsed_for_cert = {"certifications": ["Google Data Analytics"]}
final_parsed_for_cert = {"certifications": ["Google Data Analytics", "AWS Certified Developer"]}
valid_cert, cert_errors = validate_final_tailored_resume(master_parsed_for_cert, final_parsed_for_cert)
print(f"TG-04 (Invented Certification): Blocked={not valid_cert}, Errors={cert_errors}")

# TG-05: Title inflation / Seniority escalation
tg05_orig = "Software Engineering Intern at StartupX."
tg05_proposed = "Senior Software Engineer spearheading backend architecture."
tg05_scope = detect_unsupported_action_verbs_and_scope(tg05_orig, tg05_proposed)
print(f"TG-05 (Title/Scope Inflation): Blocked={bool(tg05_scope)}, Errors={tg05_scope}")

# TG-06: Valid rewrite (Wording improvement without inventing facts)
tg06_proposed = "Developed a responsive web application using React."
tg06_fab = detect_fabricated_claims(orig_bullet_1, tg06_proposed, "React developer", candidate_skills)
tg06_met = detect_unsupported_metrics(orig_bullet_1, tg06_proposed)
tg06_scope = detect_unsupported_action_verbs_and_scope(orig_bullet_1, tg06_proposed)
tg06_passed = (len(tg06_fab) == 0 and len(tg06_met) == 0 and len(tg06_scope) == 0)
print(f"TG-06 (Valid Rewrite): Passed={tg06_passed} (fab={tg06_fab}, met={tg06_met}, scope={tg06_scope})")

# TG-07: Valid reordering (Moving React higher based on JD requirements)
skills_before = ["Python", "MongoDB", "React"]
jd_text = "Looking for a Frontend Engineer with strong React and TypeScript experience."
reordered, matched, gaps, was_reordered = compute_deterministic_skill_reorder(skills_before, jd_text)
print(f"TG-07 (Valid Reordering): Reordered={reordered}, Matched={matched}, Gaps={gaps}, Changed={was_reordered}")

# TG-08: Valid project emphasis
projects = [
    ProjectEntity(id="p1", title="E-commerce Web App", technologies=["React", "Python"]),
    ProjectEntity(id="p2", title="ML Prediction Project", technologies=["Python", "Scikit-Learn", "TensorFlow"], bullets=["Trained model achieving 92% accuracy."]),
]
class MockJD:
    required_skills = ["Python", "Machine Learning", "TensorFlow"]
ranked_projects = compute_project_priorities(projects, MockJD())
print(f"TG-08 (Valid Project Emphasis): Top project={ranked_projects[0] if ranked_projects else None}, Correctly prioritized ML={ranked_projects[0] == 'ML Prediction Project'}")

# Semantic false claim vs valid rephrasing:
# Semantic test A: Candidate: "Created a REST API using Node.js" -> Tailored: "Developed RESTful APIs with Node.js"
cand_skills_node = ["Node.js", "Express"]
sem_orig_a = "Created a REST API using Node.js."
sem_prop_a = "Developed RESTful APIs with Node.js."
sem_a_fab = detect_fabricated_claims(sem_orig_a, sem_prop_a, "Node.js developer", cand_skills_node)
sem_a_scope = detect_unsupported_action_verbs_and_scope(sem_orig_a, sem_prop_a)
print(f"Semantic A (Valid rephrasing): Allowed={len(sem_a_fab) == 0 and len(sem_a_scope) == 0}")

# Semantic test B: Candidate: "Used MongoDB." -> Tailored: "Designed and optimized distributed MongoDB infrastructure at scale."
sem_orig_b = "Used MongoDB."
sem_prop_b = "Designed and optimized distributed MongoDB infrastructure at scale."
sem_b_scope = detect_unsupported_action_verbs_and_scope(sem_orig_b, sem_prop_b)
print(f"Semantic B (Scope escalation): Flagged={bool(sem_b_scope)}, Flags={sem_b_scope}")

# ============================================================
# 2. RESUME STRATEGY AUDIT (Fresher vs Experienced)
# ============================================================
print("\n--- 2. RESUME STRATEGY AUDIT ---")
# Fresher profile
fresher_profile = CandidateProfile(
    personal={"name": "Alex Fresher", "email": "alex@student.edu"},
    education=[EducationEntity(id="edu_1", institution="State Univ", degree="B.S. in Computer Science", dates="2020-2024", gpa="3.7")],
    projects=[
        ProjectEntity(id="p1", title="Campus Event App", technologies=["React", "Node.js"], bullets=["Built event registration platform."]),
        ProjectEntity(id="p2", title="Sorting Visualizer", technologies=["JavaScript", "HTML/CSS"], bullets=["Visualized algorithms."]),
    ],
    experience=[],
    skills=["JavaScript", "React", "Node.js", "Python", "Git"],
)
fresher_strat = build_resume_strategy(fresher_profile, target_role="Junior Frontend Engineer")
print(f"Fresher Strategy: CareerStage={fresher_strat.career_stage}, TemplateFamily={fresher_strat.template_family}, PageBudget={fresher_strat.page_budget}")
print(f"Fresher Section Order: {fresher_strat.section_order[:5]} (Education and Projects top-priority)")

# Experienced profile
exp_profile = CandidateProfile(
    personal={"name": "Morgan Senior", "email": "morgan@work.com"},
    education=[EducationEntity(id="edu_2", institution="Tech Institute", degree="B.S. in Software Engineering", dates="2012-2016")],
    experience=[
        WorkExperienceEntity(id="e1", company="Cloud Corp", role="Staff Engineer", dates="2020-Present", bullets=["Architected global microservices."]),
        WorkExperienceEntity(id="e2", company="Data Inc", role="Senior Software Engineer", dates="2016-2020", bullets=["Maintained distributed databases."]),
        WorkExperienceEntity(id="e3", company="Startup Z", role="Software Engineer", dates="2014-2016", bullets=["Built core API services."]),
    ],
    skills=["Go", "Kubernetes", "AWS", "Distributed Systems", "PostgreSQL", "Docker", "Terraform"],
)
exp_strat = build_resume_strategy(exp_profile, target_role="Principal Infrastructure Architect")
print(f"Experienced Strategy: CareerStage={exp_strat.career_stage}, TemplateFamily={exp_strat.template_family}, PageBudget={exp_strat.page_budget}")
print(f"Experienced Section Order: {exp_strat.section_order[:5]} (Experience dominant)")

# ============================================================
# 3. ATS ANALYSIS & SCORING SEPARATION AUDIT
# ============================================================
print("\n--- 3. ATS ANALYSIS AUDIT ---")
ats_sample_text = (
    "Jane Doe\njane@example.com • 555-0199 • San Francisco, CA\n\n"
    "PROFESSIONAL SUMMARY\nSoftware engineer specializing in React, TypeScript, and Python.\n\n"
    "TECHNICAL SKILLS\nReact, TypeScript, Python, FastAPI, PostgreSQL, Docker, Redis\n\n"
    "PROFESSIONAL EXPERIENCE\nTechCorp — Software Engineer (2022-Present)\n"
    "• Developed scalable REST APIs using FastAPI and PostgreSQL, handling 100k daily requests.\n"
    "• Optimized React application bundle size by 25% with dynamic code splitting.\n\n"
    "EDUCATION\nUC Berkeley — B.S. Computer Science (2018-2022)"
)
target_jd = "Looking for a Software Engineer with React, TypeScript, and FastAPI. Experience with PostgreSQL and Docker required."

ats_result = compute_ats_score(
    resume_text=ats_sample_text,
    jd_text=target_jd,
    parseability_score=95,
    recruiter_impact_score=80,
    skill_match_score=85,
    role_match_score=80,
)
print(f"ATS Overall Score: {ats_result.overall}/100")
print(f"ATS Categories: {[c.category_name + ': ' + str(c.points_awarded) + '/' + str(c.max_points) for c in ats_result.categories]}")
print(f"ATS Action Plan: {[a.title for a in ats_result.action_plan]}")
print(f"Knockout passed: {ats_result.knockout_passed}, Over-optimization: {ats_result.over_optimization_warning}")

# Keyword stuffing test
stuffed_text = ats_sample_text + " " + " ".join(["react"] * 50)
ats_stuffed = compute_ats_score(
    resume_text=stuffed_text,
    jd_text=target_jd,
    parseability_score=95,
    recruiter_impact_score=80,
    skill_match_score=85,
    role_match_score=80,
)
print(f"Keyword stuffing detection: Warning={ats_stuffed.over_optimization_warning}, Keyword density={ats_stuffed.keyword_density}%")

# ============================================================
# 4. EXPORT AUDIT (PDF & DOCX)
# ============================================================
print("\n--- 4. EXPORT AUDIT ---")
export_data = {
    "personal": {
        "name": "Jane Doe",
        "email": "jane@example.com",
        "phone": "555-0199",
        "location": "San Francisco, CA",
        "linkedin": "linkedin.com/in/janedoe",
    },
    "summary": "Full Stack Engineer with 3+ years experience building web applications.",
    "skills": ["Python", "React", "TypeScript", "PostgreSQL", "FastAPI", "Docker"],
    "experience_raw": [
        "TechCorp — San Francisco, CA\nSoftware Engineer (2022 - Present)\n• Developed REST APIs using FastAPI and PostgreSQL, serving 100k daily requests.\n• Optimized React frontend state management, reducing bundle size by 25%."
    ],
    "projects_raw": [
        "Task Queue\nTechnologies: Python, Redis, Docker\n• Engineered an asynchronous task queue handling 5,000 tasks/second."
    ],
    "education_raw": [
        "UC Berkeley\nB.S. in Computer Science (2018 - 2022)\nGPA: 3.8 / 4.0"
    ],
    "certifications": ["AWS Certified Solutions Architect"],
    "achievements": ["Dean's Honor List 2021"],
}

pdf_bytes = generate_pdf(export_data, candidate_name="Jane Doe", template="modern")
docx_bytes = generate_docx(export_data, candidate_name="Jane Doe", template="modern")

pdf_valid, pdf_report = verify_export_against_structured_resume(pdf_bytes, export_data, file_type="pdf")
docx_valid, docx_report = verify_export_against_structured_resume(docx_bytes, export_data, file_type="docx")

import re
print(f"PDF Export Verification: Valid={pdf_valid}, Report={pdf_report}")
print(f"DOCX Export Verification: Valid={docx_valid}, Report={docx_report}")

# Check text selectability and page count in PDF
pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
pdf_text_extracted = "".join(p.get_text() for p in pdf_doc)
print(f"PDF Page count: {pdf_doc.page_count}, Text length: {len(pdf_text_extracted)}, Selectable: {len(pdf_text_extracted) > 100}")
pdf_doc.close()

# Check DOCX openability
doc = Document(io.BytesIO(docx_bytes))
docx_text_extracted = "\n".join(p.text for p in doc.paragraphs)
print(f"DOCX Paragraph count: {len(doc.paragraphs)}, Text length: {len(docx_text_extracted)}, Openable: True")

# Consistency check between PDF and DOCX
pdf_words = set(re.findall(r"\b[a-z]{3,}\b", pdf_text_extracted.lower()))
docx_words = set(re.findall(r"\b[a-z]{3,}\b", docx_text_extracted.lower()))
common_ratio = len(pdf_words & docx_words) / max(1, len(pdf_words | docx_words))
print(f"PDF <-> DOCX Word Vocabulary Overlap: {common_ratio*100:.1f}%")

# ============================================================
# 5. USER EDIT AUDIT
# ============================================================
print("\n--- 5. USER EDIT AUDIT ---")
user_edited_good = copy.deepcopy(export_data)
user_edited_good["summary"] = "Full Stack Engineer with demonstrated background in Python and React web applications."
is_v, viols, un_tech, un_met, scope_esc = _audit_user_edits(export_data, user_edited_good, None)
print(f"User Edit (Valid rephrase): Valid={is_v}, Violations={viols}")

user_edited_bad = copy.deepcopy(export_data)
user_edited_bad["skills"].append("Kubernetes")  # Unsupported tech
user_edited_bad["experience_raw"].append("Spearheaded enterprise infrastructure overhaul saving $1,000,000.") # Unsupported metric & leadership
is_v_bad, viols_bad, un_tech_bad, un_met_bad, scope_esc_bad = _audit_user_edits(export_data, user_edited_bad, None)
print(f"User Edit (Adversarial unsupported claim): Valid={is_v_bad}, Tech={un_tech_bad}, Met={un_met_bad}, Violations={viols_bad}")

print("\n--- AUDIT COMPLETE ---")
