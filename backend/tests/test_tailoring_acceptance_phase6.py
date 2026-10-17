"""Deterministic, service-level acceptance path for resume tailoring.

The AI boundary is deterministic; upload/parsing, JD analysis, persistence,
finalization, rendering, export validation, and extraction use application code.
Timing output is diagnostic only and varies by machine.
"""

from __future__ import annotations

import io
import time
from pathlib import Path

import fitz
import pytest
from docx import Document
from mongomock_motor import AsyncMongoMockClient

from app.core.ai_service.schemas import StructuredTailoringResult
from app.core.ai_service.service import AIService
from app.core.config import Settings
from app.modules.auth import services as auth_services
from app.modules.jobs import services as jobs_services
from app.modules.resume import services as resume_services
from app.modules.tailoring import routes as tailoring_routes
from app.modules.tailoring import services as tailoring_services


@pytest.fixture
def db():
    client = AsyncMongoMockClient()
    return client["roleradar_acceptance_phase6"]


@pytest.fixture
def settings():
    return Settings(JWT_SECRET="acceptance-test-secret", EMBEDDING_PROVIDER="mock", AI_PROVIDER="mock")


@pytest.mark.asyncio
async def test_real_service_acceptance_path_upload_to_verified_exports(db, settings, monkeypatch):
    """Exercise the user journey through actual services and both export routes."""
    root = Path(__file__).parent / "fixtures"
    resume_bytes = (root / "good_resume.docx").read_bytes()
    user, _ = await auth_services.register_user(
        db, settings, "acceptance@example.com", "supersecret1", "Ananya Rao", None
    )
    user_id = str(user["_id"])
    timings: dict[str, float] = {}

    start = time.perf_counter()
    master = await resume_services.ingest_resume(
        db, settings, user_id, "good_resume.docx", resume_bytes
    )
    timings["parse_and_persist"] = time.perf_counter() - start
    assert master["parsed"]["personal"]["name"] == "Ananya Rao"
    assert "Built a Flask app" in master["raw_text"]

    jd = (
        "Backend Engineer\nCompany: Example Systems\n\n"
        "Required Skills:\n- Python\n- REST APIs\n\n"
        "Preferred Skills:\n- Docker\n"
    )
    start = time.perf_counter()
    ai_service = AIService(settings)
    job = await jobs_services.create_custom_job(
        db, "Example Systems", "Backend Engineer", jd, user_id=user_id
    )
    timings["jd_analysis_and_persistence"] = time.perf_counter() - start
    assert "Python" in job["must_have_skills"]
    assert "Docker" in job["preferred_skills"]

    async def deterministic_plan(*args, **kwargs):
        return StructuredTailoringResult()

    monkeypatch.setattr(ai_service, "generate_resume_rewrite", deterministic_plan)
    start = time.perf_counter()
    draft = await tailoring_services.generate_tailoring(db, ai_service, user_id, job_id=job["id"])
    timings["tailoring_proposal"] = time.perf_counter() - start
    assert draft["is_finalized"] is False
    assert draft["master_resume_id"] == str(master["_id"])
    assert draft["master_resume_version"] == master["version"]

    start = time.perf_counter()
    finalized = await tailoring_services.finalize_tailoring(
        db, user_id, str(draft["_id"]), settings=settings
    )
    timings["finalization"] = time.perf_counter() - start
    assert finalized["is_finalized"] is True
    assert finalized["validation_summary"]["all_checks_passed"] is True
    assert finalized["parsed"]["personal"] == master["parsed"]["personal"]

    user_ctx = {"_id": user_id, "full_name": "Ananya Rao"}
    start = time.perf_counter()
    pdf_response = await tailoring_routes.export_pdf(
        str(finalized["_id"]), template=None, current_user=user_ctx, db=db
    )
    timings["pdf_export_and_integrity_check"] = time.perf_counter() - start
    pdf = fitz.open(stream=pdf_response.body, filetype="pdf")
    pdf_text = "\n".join(page.get_text() for page in pdf)
    pdf_pages = len(pdf)
    pdf.close()

    start = time.perf_counter()
    docx_response = await tailoring_routes.export_docx(
        str(finalized["_id"]), template=None, current_user=user_ctx, db=db
    )
    timings["docx_export_and_integrity_check"] = time.perf_counter() - start
    extracted_docx = Document(io.BytesIO(docx_response.body))
    docx_text = "\n".join(p.text for p in extracted_docx.paragraphs)
    docx_text += "\n" + "\n".join(
        cell.text for table in extracted_docx.tables for row in table.rows for cell in row.cells
    )
    assert pdf_response.status_code == 200 and docx_response.status_code == 200
    for exported_text in (pdf_text, docx_text):
        assert "Ananya Rao" in exported_text
        assert "Python" in exported_text
        assert "Built a Flask app" in exported_text
    assert 1 <= pdf_pages <= finalized["resume_strategy"]["page_budget"]

    print("Acceptance stage timings (seconds):", {k: round(v, 4) for k, v in timings.items()})
    assert set(timings) == {
        "parse_and_persist",
        "jd_analysis_and_persistence",
        "tailoring_proposal",
        "finalization",
        "pdf_export_and_integrity_check",
        "docx_export_and_integrity_check",
    }
